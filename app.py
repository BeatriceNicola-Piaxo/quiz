from flask import Flask, render_template, request, jsonify, send_file
import json
import os
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import Workbook, load_workbook

BASE_DIR = Path(__file__).resolve().parent
RESULTS_FILE = BASE_DIR / "results.json"
RESULTS_XLSX_FILE = BASE_DIR / "results.xlsx"

app = Flask(__name__)


def ensure_xlsx_file():
    if not RESULTS_XLSX_FILE.exists():
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Quiz Results"
        sheet.append(["name", "score", "submitted_at"])
        workbook.save(RESULTS_XLSX_FILE)


def append_result_to_xlsx(name, score):
    ensure_xlsx_file()

    workbook = load_workbook(RESULTS_XLSX_FILE)
    sheet = workbook.active

    if sheet.max_row == 0 or sheet["A1"].value is None:
        sheet.append(["name", "score", "submitted_at"])

    if [sheet.cell(row=1, column=1).value, sheet.cell(row=1, column=2).value, sheet.cell(row=1, column=3).value] != [
        "name",
        "score",
        "submitted_at",
    ]:
        sheet.insert_rows(1)
        sheet.append(["name", "score", "submitted_at"])
        sheet.delete_rows(2, 1)

    submitted_at = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M:%S %Z")
    sheet.append([name, score, submitted_at])
    workbook.save(RESULTS_XLSX_FILE)


@app.get("/")
def index():
    return render_template("index.html")


@app.get("/health")
def health():
    return jsonify({"status": "ok"})


@app.get("/download-results")
def download_results():
    ensure_xlsx_file()
    return send_file(
        RESULTS_XLSX_FILE,
        as_attachment=True,
        download_name="quiz_results.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


@app.post("/submit")
def submit_score():
    payload = request.get_json(silent=True) or {}
    name = str(payload.get("name") or "").strip()
    score = str(payload.get("score") or "").strip()

    if not name or not score:
        return jsonify({"error": "Missing required fields: name and score"}), 400

    results = []
    if RESULTS_FILE.exists():
        try:
            with RESULTS_FILE.open("r", encoding="utf-8") as file:
                existing_results = json.load(file)
            if isinstance(existing_results, list):
                results = existing_results
        except (json.JSONDecodeError, OSError):
            results = []

    results.append({"name": name, "score": score})

    with RESULTS_FILE.open("w", encoding="utf-8") as file:
        json.dump(results, file, ensure_ascii=False, indent=2)

    append_result_to_xlsx(name, score)

    return jsonify({
        "message": "Score saved successfully",
        "name": name,
        "score": score,
    })


if __name__ == "__main__":
    app.run(host="0.0.0.0", port=int(os.environ.get("PORT", 5000)))
